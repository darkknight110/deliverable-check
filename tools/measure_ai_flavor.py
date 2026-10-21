#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""AI 味与返工量测：用真实版本链校准"什么叫改得像人写的"。

两个量：
  1. **AI 味标记**（词汇/句式/结构三层）—— 只数可数的东西，不打"像不像人"的分。
  2. **返工量** —— 相邻版本的字数变化与相似度，回答"一稿到终稿到底改了多少"。

为什么不用"盲测打分"：我们没有条件反复组织盲测；但**版本链是现成的实验记录**——
二十稿的改动方向，比一次主观打分更能说明"哪里必须改"。

用法：
    python3 tools/measure_ai_flavor.py --out docs/ai-flavor.json "某目录（md/docx）"
"""
from __future__ import annotations

import argparse
import difflib
import json
import os
import re
import statistics
import sys
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from _text import (  # noqa: E402
    cjk_chars, paragraphs, per_1k, read_text, sentences, strip_markup,
)

# ── 词汇层：最容易被一眼认出的"AI 常用词"（每个都是可数的字符串，不是主观判断）
LEXICON = {
    "管理黑话": r"赋能|抓手|闭环|生态|维度|颗粒度|底层逻辑|组合拳|护城河|心智|锚点|范式|势能|链路|沉淀",
    "连接词": r"首先|其次|再次|最后|综上|总而言之|值得注意的是|需要指出的是|不难发现",
    "对偶句式": r"不是[^，。]{2,20}而是|不仅[^，。]{2,20}而且|既[^，。]{2,20}又",
    "破折号": r"——",
    "形容词堆叠": r"(?:高效|智能|全面|深度|无缝|极致|强大|灵活|可靠|稳定){2,}",
}
CERTAIN = r"一定|必然|绝对|保证|100%|完全|彻底|零风险|无损"

# ── 结构层
PARALLEL = re.compile(r"[^，。；\n]{2,10}、[^，。；\n]{2,10}、[^，。；\n]{2,10}")
LIST_ITEM = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s+", re.M)
HEADING = re.compile(r"^#{2,6}\s+", re.M)
DIGIT = re.compile(r"\d")
LATIN = re.compile(r"[A-Za-z]{2,}")
QUOTED = re.compile(r"「[^」]{1,30}」|《[^》]{1,30}》|\"[^\"]{1,30}\"")


def extract_docx(path: str) -> str:
    import docx

    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def load_text(path: str) -> str:
    return extract_docx(path) if path.lower().endswith(".docx") else read_text(path)


def metrics(text: str) -> dict:
    md = text
    body = strip_markup(md)
    sents = sentences(body)
    paras = paragraphs(body)
    chars = cjk_chars(body)
    lexical = {k: len(re.findall(p, body)) for k, p in LEXICON.items()}
    # 抽象句：既没有数字、也没有专名/英文/引号内容 —— 说完等于没说的高发形态
    abstract = [s for s in sents if not (DIGIT.search(s) or LATIN.search(s) or QUOTED.search(s))]
    return {
        "chars": chars,
        "sentences": len(sents),
        "paragraphs": len(paras),
        "avg_sentence_len": round(statistics.mean(len(s) for s in sents), 1) if sents else 0,
        "avg_paragraph_len": round(statistics.mean(len(p) for p in paras), 1) if paras else 0,
        "headings_per_1k": per_1k(len(HEADING.findall(body)), chars),
        "list_items": len(LIST_ITEM.findall(body)),
        "list_ratio": round(len(LIST_ITEM.findall(body)) / max(1, len(paras)), 2),
        "parallel_triples": len(PARALLEL.findall(body)),
        "parallel_per_1k": per_1k(len(PARALLEL.findall(body)), chars),
        "lexical": lexical,
        "lexical_per_1k": {k: per_1k(v, chars) for k, v in lexical.items()},
        "certain_per_1k": per_1k(len(re.findall(CERTAIN, body)), chars),
        "abstract_ratio": round(len(abstract) / max(1, len(sents)), 3),
    }


VERSION = re.compile(r"_V(\d+)(?:\.(\d+))?")
DOC_TYPE = re.compile(r"(设计方案|施工手册|构想稿)")


def version_key(name: str) -> tuple | None:
    m = VERSION.search(name)
    if not m:
        return None
    return (int(m.group(1)), int(m.group(2) or 0))


def doc_family(name: str) -> str:
    m = DOC_TYPE.search(name)
    return m.group(1) if m else "其他"


def chain_analysis(docs: list[dict]) -> list[dict]:
    """同一文档类型按版本排序，算相邻两版的相似度与改动量。"""
    groups = {}
    for d in docs:
        groups.setdefault(d["family"], []).append(d)
    chains = []
    for family, items in groups.items():
        items.sort(key=lambda x: x["version"] or (0, 0))
        steps = []
        for prev, cur in zip(items, items[1:]):
            a, b = prev["text"], cur["text"]
            ratio = round(difflib.SequenceMatcher(None, a, b, autojunk=False).ratio(), 3)
            steps.append({
                "from": prev["name"],
                "to": cur["name"],
                "chars_from": cjk_chars(a),
                "chars_to": cjk_chars(b),
                "delta_chars": cjk_chars(b) - cjk_chars(a),
                "similarity": ratio,
                "rewrite_share": round(1 - ratio, 3),
            })
        chains.append({"family": family, "versions": [i["name"] for i in items], "steps": steps})
    return chains


def _redact_path(path: str) -> str:
    import hashlib

    return "doc-" + hashlib.sha1(path.encode("utf-8")).hexdigest()[:10]


def main() -> int:
    ap = argparse.ArgumentParser(description="量测 AI 味标记与版本返工量")
    ap.add_argument("inputs", nargs="+", help="目录或 .md/.docx 文件")
    ap.add_argument("--out", help="输出 JSON")
    ap.add_argument("--json-only", action="store_true")
    ap.add_argument("--redact", action="store_true", help="输出里隐去真实路径（公开前必须加）")
    args = ap.parse_args()

    paths = []
    for item in args.inputs:
        if os.path.isdir(item):
            for root, _dirs, files in os.walk(item):
                paths += [os.path.join(root, f) for f in files if f.lower().endswith((".md", ".docx"))]
        else:
            paths.append(item)
    paths = sorted(set(p for p in paths if not os.path.basename(p).startswith("~$")))

    docs = []
    for p in paths:
        name = os.path.basename(p)
        try:
            text = load_text(p)
            m = metrics(text)
        except Exception as exc:
            print(f"! 跳过 {name}: {exc}", file=sys.stderr)
            continue
        m.update({
            "path": _redact_path(p) if args.redact else p,
            "name": name,
            "version": version_key(name),
            "family": doc_family(name),
            "text": text,
        })
        docs.append(m)

    chains = chain_analysis(docs)
    payload = {
        "documents": [{k: v for k, v in d.items() if k != "text"} for d in docs],
        "chains": chains,
    }
    text_out = json.dumps(payload, ensure_ascii=False, indent=2)
    if args.out:
        os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
        with open(args.out, "w", encoding="utf-8") as fh:
            fh.write(text_out + "\n")
    if args.json_only:
        print(text_out)
    else:
        print(f"{len(docs)} 份文档")
        for ch in chains:
            if len(ch["versions"]) < 2:
                continue
            print(f"\n版本链：{ch['family']}（{len(ch['versions'])} 版）")
            for s in ch["steps"]:
                print(f"  {s['chars_from']:>6} → {s['chars_to']:>6}  "
                      f"改动 {s['rewrite_share']:.0%}  {s['from'][:22]} → {s['to'][:22]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
