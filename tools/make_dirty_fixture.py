#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成一张"看起来正常、算出来错 176%"的脏表，用于验证 audit_data.py 真的能拦住。

为什么要把脏表做成脚本而不是提交一个二进制：
  1. 差异可读：任何人能一眼看出"哪里脏"；
  2. 可复跑：CI 里现场生成，不依赖仓库里的二进制；
  3. 可回归：改成更脏/更干净，看检查项是否按预期变化。

用法：
    python3 tools/make_dirty_fixture.py --out /tmp/dirty.xlsx
"""
from __future__ import annotations

import argparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font


def build(path: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "订单明细"

    # ── 两级表头 + 合并单元格（读表时列名会变成 Unnamed）
    ws["A1"] = "订单"
    ws["B1"] = "区域"
    ws["C1"] = "金额"
    ws["D1"] = "状态"
    ws.merge_cells("C1:D1")
    ws["C2"] = "（元）"
    ws["D2"] = "（当前）"
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:B2")
    for cell in ("A1", "B1", "C1", "C2", "D2"):
        ws[cell].font = Font(bold=True)
        ws[cell].alignment = Alignment(horizontal="center")

    # ── 明细：金额写成"带千分位的文本"
    rows = [
        ("SO-1001", "华东", "12,000", "已发货"),
        ("SO-1002", "华东", "8,500.50", "已发货"),
        ("华东小计", "", "20,500.50", ""),          # ← 小计混在明细里
        ("SO-1003", "华南", "3,300", "待发货"),
        ("SO-1003", "华南", "3,300", "待发货"),      # ← 主键重复
        ("合计", "", "27,100.50", ""),              # ← 合计也在明细里
    ]
    for i, r in enumerate(rows, start=3):
        for j, v in enumerate(r, start=1):
            ws.cell(row=i, column=j, value=v)

    # ── 单位写进数值的示例列
    ws["E2"] = "重量"
    ws["E3"] = "1.2 吨"
    ws["E4"] = "0.8 吨"

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 14
    wb.save(path)
    return path


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="/tmp/dirty-orders.xlsx")
    args = ap.parse_args()
    out = build(args.out)
    print(out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
