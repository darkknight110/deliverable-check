#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Word 表格的数值审计。"""
import os
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

import docx  # noqa: E402

from deliverable_check import audit_data, preflight  # noqa: E402


def make(rows):
    d = docx.Document()
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    td = tempfile.mkdtemp()
    p = os.path.join(td, "表.docx")
    d.save(p)
    return p


class TestDocxTables(unittest.TestCase):
    def test_合计混进明细_阻断(self):
        p = make([["项目", "金额"], ["A", "1200"], ["B", "800"], ["小计", "2000"]])
        r = audit_data.audit_docx_tables(p)
        self.assertEqual(r["verdict"], "blocked")
        self.assertIn("docx-total-in-body", {f["id"] for f in r["findings"]})

    def test_干净表_不误报(self):
        p = make([["项目", "金额"], ["A", "1200"], ["B", "800"]])
        r = audit_data.audit_docx_tables(p)
        self.assertEqual(r["verdict"], "clean", r)

    def test_preflight_汇总里能看到表格结论(self):
        p = make([["项目", "金额"], ["A", "1200"], ["小计", "1200"]])
        rep = preflight.run([os.path.dirname(p)])
        self.assertEqual(rep["verdict"], "blocked")
        self.assertTrue(rep["docx_tables"][0]["blocking"])


if __name__ == "__main__":
    unittest.main()
